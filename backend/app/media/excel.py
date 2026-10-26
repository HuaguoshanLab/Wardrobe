import io
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath
from typing import Protocol, cast

from defusedxml.ElementTree import fromstring
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, TwoCellAnchor
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.errors import AppError
from app.domain.taxonomy import MULTI, TAXONOMY
from app.media.images import demo_image

COLUMNS = [
    "图片",
    "品类",
    "名称",
    "颜色",
    "季节",
    "场合",
    "风格",
    "材质",
    "厚度",
    "保暖程度",
    "自定义标签",
    "状态",
]
FIELDS = [
    "image",
    "category",
    "name",
    "color_codes",
    "seasons",
    "occasion_codes",
    "style_codes",
    "material",
    "thickness",
    "warmth",
    "tags",
    "availability",
]


class EmbeddedImage(Protocol):
    anchor: OneCellAnchor | TwoCellAnchor

    def _data(self) -> bytes: ...


class ImageSheet(Protocol):
    _images: list[EmbeddedImage]


@dataclass
class ParsedRow:
    number: int
    attributes: dict[str, object] = field(default_factory=dict)
    tags: list[str] = field(default_factory=list)
    image: bytes | None = None
    errors: list[str] = field(default_factory=list)


def validate_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(x.file_size for x in entries) > 250 * 1024 * 1024:
                raise ValueError("archive limit")
            if len({x.filename for x in entries}) != len(entries):
                raise ValueError("duplicate archive members")
            for entry in entries:
                path = PurePosixPath(entry.filename)
                if path.is_absolute() or ".." in path.parts or "\\" in entry.filename:
                    raise ValueError("invalid path")
                if entry.file_size > max(entry.compress_size, 1) * 500:
                    raise ValueError("compression ratio")
                name = entry.filename.lower()
                if any(
                    x in name for x in ("vbaproject", "externallinks/", "cellimages", "richdata/")
                ):
                    raise ValueError("unsupported image/macro/link structure")
                if name.endswith(".rels"):
                    root = fromstring(archive.read(entry))
                    if any(x.get("TargetMode") == "External" for x in root):
                        raise ValueError("external links")
    except Exception as exc:
        raise AppError(
            "INVALID_WORKBOOK", "Excel 不安全、超限或包含不支持的宏/外链/单元格图片", 422
        ) from exc


def parse_excel(data: bytes) -> list[ParsedRow]:
    validate_archive(data)
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
        if "meta" not in workbook or "garments" not in workbook:
            raise ValueError("missing sheets")
        meta = workbook["meta"]
        if meta["A1"].value != "template_version" or meta["B1"].value != 1:
            raise ValueError("template version")
        sheet = cast(Worksheet, workbook["garments"])
        if (
            [sheet.cell(1, n).value for n in range(1, 13)] != COLUMNS
            or sheet.max_row > 10000
            or sheet.max_column > 64
        ):
            raise ValueError("headers or row limit")
        rows: dict[int, ParsedRow] = {}
        for number, values in enumerate(sheet.iter_rows(min_row=2), start=2):
            if not any(cell.value is not None for cell in values):
                continue
            row = rows.setdefault(number, ParsedRow(number))
            for index, cell in enumerate(values):
                if cell.value is None:
                    continue
                if cell.data_type == "f" or index >= len(FIELDS):
                    row.errors.append("不支持公式或额外数据列")
                    continue
                name = FIELDS[index]
                raw = str(cell.value).strip()
                if name == "image":
                    row.errors.append("图片列应插入嵌入图片，不填写文字或 URL")
                elif name == "tags":
                    row.tags = [v.strip() for v in raw.split(";") if v.strip()]
                elif name == "name":
                    row.attributes[name] = raw
                else:
                    lookup = {v: k for k, v in TAXONOMY[name].items()}
                    parts = [v.strip() for v in raw.split(";") if v.strip()]
                    if name == "seasons" and parts == ["四季"]:
                        parts = list(TAXONOMY[name])
                    converted = [lookup.get(v, v) for v in parts]
                    row.attributes[name] = (
                        converted if name in MULTI else converted[0] if converted else None
                    )
        for drawing in cast(ImageSheet, sheet)._images:
            anchor = drawing.anchor
            if not hasattr(anchor, "_from"):
                raise ValueError("absolute image anchor")
            number = anchor._from.row + 1
            if number < 2:
                raise ValueError("image in header")
            row = rows.setdefault(number, ParsedRow(number))
            if anchor._from.col != 0 or row.image is not None:
                row.errors.append("图片必须位于图片列，且每行仅一张")
            elif hasattr(anchor, "to") and (
                anchor.to.row > anchor._from.row + 1
                or (anchor.to.row == anchor._from.row + 1 and anchor.to.rowOff > 0)
            ):
                row.errors.append("图片跨行，无法确定所属行")
            elif hasattr(anchor, "ext") and (
                anchor._from.rowOff + anchor.ext.cy
                > (sheet.row_dimensions[number].height or sheet.sheet_format.defaultRowHeight or 15)
                * 12700
                + 12700
            ):
                row.errors.append("图片跨行，请增大该行行高后重新导入")
            else:
                row.image = drawing._data()
        if not rows or len(rows) > 100:
            raise ValueError("empty or >100 rows")
        for row in rows.values():
            if row.image is None:
                row.errors.append("缺少图片，请修正 Excel 后重新导入")
        return sorted(rows.values(), key=lambda row: row.number)
    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            "INVALID_TEMPLATE", "请使用模板 v1，最多 100 行，图片需嵌入并锚定到对应行", 422
        ) from exc


def template() -> bytes:
    book = Workbook()
    sheet = cast(Worksheet, book.active)
    sheet.title = "garments"
    sheet.append(COLUMNS)
    sheet.append(
        [
            None,
            "上装",
            "示例：请替换图片与属性",
            "绿色",
            "春;秋",
            "日常",
            "休闲",
            None,
            None,
            None,
            None,
            "可穿",
        ]
    )
    sheet.row_dimensions[2].height = 120
    image = ExcelImage(io.BytesIO(demo_image()))
    image.width = image.height = 140
    sheet.add_image(image, "A2")
    sheet.freeze_panes = "C2"
    for index, cell in enumerate(sheet[1], start=1):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="435E44")
        sheet.column_dimensions[get_column_letter(index)].width = 22
    meta = book.create_sheet("meta")
    meta.append(["template_version", 1])
    help_sheet = book.create_sheet("说明与词表")
    help_sheet.append(["使用说明", "模板含一行示例，请替换或删除；图片和品类必填。"])
    help_sheet.append(
        ["图片", "插入浮动嵌入图片（不选置于单元格），左上角置于图片列对应行，一行一图。"]
    )
    help_sheet.append(["多选", "用英文分号 ; 分隔；空白为未知；不要使用公式或外部链接。"])
    for name, choices in TAXONOMY.items():
        help_sheet.append([name, ";".join(choices.values())])
    help_sheet.column_dimensions["A"].width = 24
    help_sheet.column_dimensions["B"].width = 100
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()
