import io
import zipfile

import pytest
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from app.core.errors import AppError
from app.media.excel import parse_excel, template
from app.media.images import cutout, demo_image, normalize


def workbook_bytes(book):
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def test_orientation_and_metadata_stripping():
    source = Image.new("RGB", (60, 30), "green")
    exif = Image.Exif()
    exif[274] = 6
    exif[270] = "private data"
    stream = io.BytesIO()
    source.save(stream, format="JPEG", exif=exif)
    data, width, height = normalize(stream.getvalue())
    assert (width, height) == (30, 60)
    assert not Image.open(io.BytesIO(data)).getexif()


def test_heic_conversion_and_real_cutout():
    source = Image.open(io.BytesIO(demo_image()))
    heic = io.BytesIO()
    source.save(heic, format="HEIF")
    normalized, width, height = normalize(heic.getvalue())
    assert (width, height) == (480, 480)
    result = Image.open(io.BytesIO(cutout(normalized)))
    assert result.getchannel("A").getextrema() == (0, 255)


@pytest.mark.parametrize("source", [b"not an image", b"<svg></svg>", b"GIF89a"])
def test_invalid_image_rejected(source):
    with pytest.raises(AppError, match="图片"):
        normalize(source)


def test_image_byte_limit():
    with pytest.raises(AppError) as exc:
        normalize(b"a" * (20 * 1024 * 1024 + 1))
    assert exc.value.code == "IMAGE_TOO_LARGE"


def test_template_roundtrip_preserves_attributes():
    rows = parse_excel(template())
    assert len(rows) == 1 and rows[0].image and not rows[0].errors
    assert rows[0].attributes["category"] == "top"
    assert rows[0].attributes["seasons"] == ["spring", "autumn"]


def test_excel_anchors_not_zip_image_order_and_row_errors():
    book = load_workbook(io.BytesIO(template()))
    sheet = book["garments"]
    sheet["B4"] = "鞋"
    sheet["B3"] = "错误品类"
    sheet["D3"] = "=1+1"
    sheet.row_dimensions[4].height = 120
    image = ExcelImage(io.BytesIO(demo_image()))
    image.width = image.height = 120
    sheet.add_image(image, "A4")
    sheet._images.reverse()
    rows = parse_excel(workbook_bytes(book))
    assert [row.number for row in rows] == [2, 3, 4]
    assert rows[2].image and rows[2].attributes["category"] == "shoes"
    assert rows[1].image is None and len(rows[1].errors) == 2


def test_excel_cross_row_and_multiple_images_rejected():
    book = load_workbook(io.BytesIO(template()))
    book["garments"].row_dimensions[2].height = 15
    assert "跨行" in ";".join(parse_excel(workbook_bytes(book))[0].errors)
    book = load_workbook(io.BytesIO(template()))
    image = ExcelImage(io.BytesIO(demo_image()))
    book["garments"].add_image(image, "A2")
    assert "每行仅一张" in ";".join(parse_excel(workbook_bytes(book))[0].errors)


@pytest.mark.parametrize("filename", ["../escape", "xl/vbaProject.bin", "xl/richData/a.xml"])
def test_unsafe_workbook_rejected(filename):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(filename, "data")
    with pytest.raises(AppError) as exc:
        parse_excel(buffer.getvalue())
    assert exc.value.code == "INVALID_WORKBOOK"
